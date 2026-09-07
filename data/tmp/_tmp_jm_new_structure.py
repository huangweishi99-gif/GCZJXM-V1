# -*- coding: utf-8 -*-
import io, sys
sys.path.insert(0, r'd:\projects\cost-estimation')
import openpyxl
NEW = r'd:\projects\cost-estimation\data\exports\江门新会骏凯豪庭五期三标项目1栋05户型精装修工程-投标报价_成本填写.xlsx'
wb = openpyxl.load_workbook(NEW, data_only=True, read_only=True)
out = io.open(r'd:\projects\cost-estimation\data\tmp\_tmp_jm_new_structure_out.txt', 'w', encoding='utf-8')
for ws in wb.worksheets:
    out.write(f'--- sheet: {ws.title}  max_row={ws.max_row} max_col={ws.max_column}\n')
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        vals = ['' if v is None else str(v).replace('\n', '|')[:18] for v in row[:20]]
        out.write(f'  r{i}: {vals}\n')
    out.write('\n')
wb.close()
out.close()
print('done')
