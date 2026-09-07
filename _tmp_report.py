# -*- coding: utf-8 -*-
"""临时：生成完整控制价对比分析报告（按原因分类，含建议新成本），用后即删。"""
import io, sys, os, csv
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import win32com.client
from openpyxl.utils import get_column_letter

XLS_A = r"清单数据资料\甲方招标清单\云湾花园项目-5#6#户内、公区，样板，架空层装修工程专业分包-投标清单-20260825.xls"
XLS_B = r"清单数据资料\甲方招标清单\云湾花园项目-会所售楼部装修工程专业分包-投标清单-20260825.xls"
XLSX = r"清单数据资料\甲方招标清单\会所、5#6#楼公区及户内、样板间、架空层安装工程投标清单-20260825.xlsx"
OUT = r"data\tmp\重填分析报告.csv"
OUT_MD = r"data\exports\云湾控制价对比分析报告.md"

recs = []


def classify(diff, cost, ctrl):
    if cost == 0 and ctrl > 0:
        return "整项漏填(成本0)"
    if diff is None:
        return "数据异常"
    if diff > 10:
        return "成本高估"
    if diff < -10:
        return "成本低估(疑漏主材)"
    return "合理(±10%)"


xl = win32com.client.DispatchEx("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    for tag, path in (("5#6#", XLS_A), ("会所xls", XLS_B)):
        wb = xl.Workbooks.Open(os.path.abspath(path), UpdateLinks=0, ReadOnly=True)
        ws = wb.Worksheets("清单表")
        ur = ws.UsedRange
        n_r = ur.Rows.Count
        val = ur.Value
        for i in range(n_r):
            r = ur.Row + i
            row = val[i]
            if not isinstance(row[0], (int, float)):
                continue
            name = str(row[1] or "")[:30]
            unit = row[5] or ""
            qty = row[6] or 0
            cost = row[19] or 0   # T
            ctrl = row[16] or 0   # Q
            u = row[20] or 0      # 主材
            v = row[21] or 0      # 损耗率
            w = row[22] or 0      # 人工
            x = row[23] or 0      # 辅材
            y = row[24] or 0      # 机械
            if not all(isinstance(z, (int, float)) for z in (cost, ctrl, u, w, x, y)):
                cost = ctrl = u = w = x = y = 0
            cost *= 1.0; ctrl *= 1.0
            ratio = (cost * 1.12 / ctrl) if ctrl else None
            diff = round((ratio - 1) * 100, 1) if ratio else None
            cat = classify(diff, cost, ctrl)
            newcost = round(ctrl / 1.12, 2) if ctrl else 0
            recs.append({
                "文件": tag, "sheet": "清单表", "行": r, "项目名称": name, "单位": str(unit),
                "工程量": qty, "成本单价": cost, "控制价": ctrl,
                "成本x1.12": round(cost * 1.12, 2), "差异%": diff, "原因分类": cat,
                "建议新成本": newcost,
                "原主材": round(u, 2), "损耗率": v, "原人工": round(w, 2),
                "原辅材": round(x, 2), "原机械": round(y, 2),
            })
        wb.Close(False)
finally:
    xl.Quit()

import openpyxl
wbv = openpyxl.load_workbook(XLSX, data_only=True)
for sn in ("会所售楼处消防电工程", "会所售楼处消防水工程"):
    ws = wbv[sn]
    for r in range(5, ws.max_row + 1):
        if not isinstance(ws.cell(row=r, column=1).value, (int, float)):
            continue
        name = str(ws.cell(row=r, column=3).value or "")[:30]
        unit = ws.cell(row=r, column=5).value or ""
        qty = ws.cell(row=r, column=6).value or 0
        cost = ws.cell(row=r, column=17).value or 0   # Q
        ctrl = ws.cell(row=r, column=14).value or 0   # N
        u = ws.cell(row=r, column=18).value or 0      # R 主材
        v = ws.cell(row=r, column=19).value or 0      # S 损耗率
        w = ws.cell(row=r, column=20).value or 0      # T 人工
        x = ws.cell(row=r, column=21).value or 0      # U 辅材
        y = ws.cell(row=r, column=22).value or 0      # V 机械
        if not all(isinstance(z, (int, float)) for z in (cost, ctrl, u, w, x, y)):
            cost = ctrl = u = w = x = y = 0
        cost *= 1.0; ctrl *= 1.0
        ratio = (cost * 1.12 / ctrl) if ctrl else None
        diff = round((ratio - 1) * 100, 1) if ratio else None
        cat = classify(diff, cost, ctrl)
        newcost = round(ctrl / 1.12, 2) if ctrl else 0
        recs.append({
            "文件": "会所安装xlsx", "sheet": sn, "行": r, "项目名称": name, "单位": str(unit),
            "工程量": qty, "成本单价": cost, "控制价": ctrl,
            "成本x1.12": round(cost * 1.12, 2), "差异%": diff, "原因分类": cat,
            "建议新成本": newcost,
            "原主材": round(u, 2), "损耗率": v, "原人工": round(w, 2),
            "原辅材": round(x, 2), "原机械": round(y, 2),
        })
wbv.close()

with open(OUT, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(recs[0].keys()))
    w.writeheader()
    for r in recs:
        w.writerow(r)

# 汇总
from collections import Counter
cat_count = Counter(r["原因分类"] for r in recs)
print("分类统计:")
for k, v in cat_count.most_common():
    print(f"  {k}: {v}")

for tag in ("5#6#", "会所xls", "会所安装xlsx"):
    sub = [r for r in recs if r["文件"] == tag]
    cc = Counter(r["原因分类"] for r in sub)
    print(f"\n[{tag}] 共{len(sub)}行:", dict(cc))

# 建议新成本 vs 控制价 一致性样例
print("\n===== 重填效果样例（前 10 大金额影响） =====")
for r in sorted(recs, key=lambda x: abs((x["控制价"] or 0) - (x["成本单价"] or 0)) * (x["工程量"] or 0), reverse=True)[:10]:
    print(f"  {r['文件'][:8]:8} {r['sheet'][:10]:10} 行{r['行']:>4} | {r['项目名称']:<22} | 成本{r['成本单价']:>9} → {r['建议新成本']:>9} | 控制价{r['控制价']:>9} | {r['原因分类']}")

print(f"\n报告已输出: {OUT}")
