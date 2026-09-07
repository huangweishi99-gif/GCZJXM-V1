# -*- coding: utf-8 -*-
"""临时：读取三份文件完整表头（行1-4）全部列，用后即删。"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import win32com.client
from openpyxl.utils import get_column_letter

XLS_A = r"清单数据资料\甲方招标清单\云湾花园项目-5#6#户内、公区，样板，架空层装修工程专业分包-投标清单-20260825.xls"
XLS_B = r"清单数据资料\甲方招标清单\云湾花园项目-会所售楼部装修工程专业分包-投标清单-20260825.xls"
XLSX = r"清单数据资料\甲方招标清单\会所、5#6#楼公区及户内、样板间、架空层安装工程投标清单-20260825.xlsx"

xl = win32com.client.DispatchEx("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    for tag, path in (("5#6#", XLS_A), ("会所xls", XLS_B)):
        print("\n##########", tag)
        wb = xl.Workbooks.Open(os.path.abspath(path), UpdateLinks=0, ReadOnly=True)
        ws = wb.Worksheets("清单表")
        for r in range(1, 5):
            vals = []
            for c in range(1, 40):
                v = ws.Cells(r, c).Value
                if v is not None:
                    s = str(v).replace("\n", "⏎")[:40]
                    vals.append(f"{get_column_letter(c)}{r}={s}")
            print(" | ".join(vals))
        wb.Close(False)
finally:
    xl.Quit()

print("\n########## xlsx 消防电/消防水 表头")
import openpyxl
wbv = openpyxl.load_workbook(XLSX, data_only=True)
wbf = openpyxl.load_workbook(XLSX, data_only=False)
for sn in ("会所售楼处消防电工程", "会所售楼处消防水工程"):
    wsv, wsf = wbv[sn], wbf[sn]
    print(f"\n## {sn} max_row={wsv.max_row} max_col={wsv.max_column}")
    for r in range(1, 5):
        vals = []
        for c in range(1, wsv.max_column + 1):
            v = wsv.cell(row=r, column=c).value
            if v is not None:
                s = str(v).replace("\n", "⏎")[:40]
                vals.append(f"{get_column_letter(c)}{r}={s}")
        print(" | ".join(vals))
    # 数据行样本 5-8
    for r in range(5, 9):
        vals = []
        for c in range(1, wsv.max_column + 1):
            v = wsv.cell(row=r, column=c).value
            if v is not None:
                s = str(v).replace("\n", "⏎")[:30]
                vals.append(f"{get_column_letter(c)}{r}={s}")
        print(" | ".join(vals))
wbv.close()
wbf.close()
