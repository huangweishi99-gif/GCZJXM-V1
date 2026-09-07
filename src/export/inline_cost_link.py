# -*- coding: utf-8 -*-
"""
清单成本链接交付入口。

读甲方招标表 → 智能识别连接画像（清单询 / 其中分解 / 国标简表）
→ 在甲方末列后追加统一成本分析区，并按画像桥接报价列。
实现见 ``src/export/cost_analysis_template.py``。
"""
from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

from src.db.database import resolve_db_path
from src.ingest.detector import detect_sheet_layout, is_pricing_row
from src.ingest.formula_eval import eval_cell
from src.knowledge.query import PricingContext
from src.knowledge.repository import KnowledgeRepository
from src.link.dedupe import MasterRowRef, make_dedupe_key
from src.pricing.component_judge import judge_line_components
from src.pricing.engine import PricingEngine

SKIP_SHEET_KEYWORDS = ("说明", "汇总", "材料表", "文本")


@dataclass
class CostParts:
    material_main: float = 0.0
    material_loss_rate: float = 0.0
    labor: float = 0.0
    material_aux: float = 0.0
    machinery: float = 0.0
    source: str = ""

    def usable(self) -> bool:
        return any(
            v > 0
            for v in (
                self.material_main,
                self.labor,
                self.material_aux,
                self.machinery,
            )
        )


@dataclass
class InlineLinkReport:
    output_path: str
    sheets: List[Dict[str, Any]] = field(default_factory=list)
    reference_hits: int = 0
    judged: int = 0
    masters: int = 0
    slaves: int = 0
    cross_sheet_slaves: int = 0


def _cell_num(val) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, str) and val.strip().startswith("="):
        return None
    try:
        return float(str(val).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None


def _row_text(row: Tuple[Any, ...], col: Optional[int]) -> str:
    if col is None or col >= len(row):
        return ""
    v = row[col]
    return "" if v is None else str(v).strip()


def _cost_indices(column_map: Dict[str, int]) -> Dict[str, Optional[int]]:
    """读取成本区列（0-based），优先 _cost_side。"""
    side = bool(column_map.get("_cost_side"))
    if side:
        return {
            "main": column_map.get("material_main_cost"),
            "loss": column_map.get("material_loss_rate_cost")
            or column_map.get("material_loss_rate"),
            "labor": column_map.get("labor_cost"),
            "aux": column_map.get("material_aux_cost"),
            "mach": column_map.get("machinery_cost"),
        }
    return {
        "main": column_map.get("material_main"),
        "loss": column_map.get("material_loss_rate"),
        "labor": column_map.get("labor"),
        "aux": column_map.get("material_aux"),
        "mach": column_map.get("machinery"),
    }


def _col_num(
    row: Tuple[Any, ...],
    col_idx: Optional[int],
    *,
    ws=None,
    excel_row: Optional[int] = None,
    cache: Optional[dict] = None,
    wb=None,
) -> Optional[float]:
    if col_idx is None:
        return None
    if col_idx < len(row):
        v = _cell_num(row[col_idx])
        if v is not None:
            return v
    if ws is None or excel_row is None:
        return None
    return eval_cell(ws, excel_row, col_idx + 1, cache, wb=wb)


def _read_row_cost(
    row: Tuple[Any, ...],
    column_map: Dict[str, int],
    *,
    ws=None,
    excel_row: Optional[int] = None,
    cache: Optional[dict] = None,
    wb=None,
) -> Optional[CostParts]:
    """读成本人材机；成本区为空时回退左侧「其中」桥接列（含跨表公式）。"""
    idx = _cost_indices(column_map)

    def take(field: str) -> Optional[float]:
        return _col_num(
            row, idx.get(field), ws=ws, excel_row=excel_row, cache=cache, wb=wb
        )

    main = take("main")
    labor = take("labor")
    aux = take("aux")
    mach = take("mach")
    loss = take("loss")
    cp = CostParts(
        material_main=float(main or 0),
        material_loss_rate=float(loss or 0),
        labor=float(labor or 0),
        material_aux=float(aux or 0),
        machinery=float(mach or 0),
    )
    if cp.usable():
        return cp

    # 链接行：左侧人工 / 主材(已含损耗) / 辅材及机械
    left_l = _col_num(
        row, column_map.get("labor"), ws=ws, excel_row=excel_row, cache=cache, wb=wb
    )
    left_m = _col_num(
        row,
        column_map.get("material_main"),
        ws=ws,
        excel_row=excel_row,
        cache=cache,
        wb=wb,
    )
    left_h = _col_num(
        row, column_map.get("machinery"), ws=ws, excel_row=excel_row, cache=cache, wb=wb
    )
    if any(x is not None and x > 0 for x in (left_l, left_m, left_h)):
        return CostParts(
            material_main=float(left_m or 0),
            material_loss_rate=0.0,
            labor=float(left_l or 0),
            material_aux=0.0,
            machinery=float(left_h or 0),
        )

    # 仅有成本单价（细分空）：整价落主材，保证引用可套用
    cost_up = _col_num(
        row,
        column_map.get("cost_unit_price"),
        ws=ws,
        excel_row=excel_row,
        cache=cache,
        wb=wb,
    )
    if cost_up is None or cost_up <= 0:
        cost_up = _col_num(
            row,
            column_map.get("unit_price"),
            ws=ws,
            excel_row=excel_row,
            cache=cache,
            wb=wb,
        )
    if cost_up and cost_up > 0:
        return CostParts(material_main=float(cost_up))
    return None


# 同名单价离散极大，禁止 (名称,单位) 弱回退
_NO_NAME_UNIT_FALLBACK = frozenset(
    {"低压电气开关", "装饰灯", "装饰风口", "灯具", "开关", "插座"}
)


def build_reference_price_index(
    reference_paths: List[str | Path],
) -> Tuple[Dict[str, CostParts], Dict[Tuple[str, str], CostParts], Dict[str, CostParts]]:
    by_key: Dict[str, CostParts] = {}
    by_name_unit: Dict[Tuple[str, str], CostParts] = {}
    by_cal: Dict[str, CostParts] = {}

    for path in reference_paths:
        fp = Path(path)
        if not fp.exists():
            continue
        src_label = fp.stem
        # .xls 先转为临时 xlsx，openpyxl 才能读公式
        work = fp
        tmp_xlsx: Optional[Path] = None
        if fp.suffix.lower() == ".xls":
            tmp_xlsx = (
                resolve_db_path(None).parent
                / "exports"
                / f"_ref_{fp.stem}.xlsx"
            )
            _materialize_workbook_copy(fp, tmp_xlsx)
            work = tmp_xlsx
        # data_only=False：跨表链接/公式可求值，避免空缓存丢价
        wb = load_workbook(work, data_only=False)
        for sn in wb.sheetnames:
            if any(k in sn for k in SKIP_SHEET_KEYWORDS):
                continue
            df = pd.read_excel(work, sheet_name=sn, header=None)
            layout = detect_sheet_layout(df, sn)
            if not layout:
                continue
            cm = layout.column_map
            if not cm.get("_cost_side") and "labor" not in cm:
                continue
            name_i = cm.get("name")
            feat_i = cm.get("feature")
            unit_i = cm.get("unit")
            if name_i is None or unit_i is None:
                continue
            ws = wb[sn]
            cache: dict = {}
            for ri in range(layout.data_start_row, len(df)):
                row = tuple(df.iloc[ri])
                if not is_pricing_row(row, cm):
                    continue
                excel_row = ri + 1
                cost = _read_row_cost(
                    row, cm, ws=ws, excel_row=excel_row, cache=cache, wb=wb
                )
                if not cost or not cost.usable():
                    continue
                name = _row_text(row, name_i)
                feature = _row_text(row, feat_i)
                unit = _row_text(row, unit_i)
                nn, un, sig = make_dedupe_key(name, feature, unit)
                full = f"{nn}|{sig}|{un}"
                cost.source = f"ref:{src_label}/{sn}"
                if full not in by_key:
                    by_key[full] = cost
                nu = (nn, un)
                if nn not in _NO_NAME_UNIT_FALLBACK and nu not in by_name_unit:
                    by_name_unit[nu] = cost
                from src.knowledge.calibration_match import make_calibration_match_key

                for pid in ("shenhai_dasha", ""):
                    ck = make_calibration_match_key(
                        name, feature, unit, project_id=pid
                    )
                    # vent_band 按宽度聚合，交付时不能把不同长度套成同价
                    if not ck or ck.startswith("vent_band|"):
                        continue
                    if ck not in by_cal:
                        by_cal[ck] = cost
        wb.close()
    return by_key, by_name_unit, by_cal


def _lookup_reference(
    name: str,
    feature: str,
    unit: str,
    by_key: Dict[str, CostParts],
    by_name_unit: Dict[Tuple[str, str], CostParts],
    by_cal: Optional[Dict[str, CostParts]] = None,
) -> Optional[CostParts]:
    nn, un, sig = make_dedupe_key(name, feature, unit)
    full = f"{nn}|{sig}|{un}"
    if full in by_key:
        return by_key[full]
    if by_cal:
        from src.knowledge.calibration_match import make_calibration_match_key

        for pid in ("shenhai_dasha", ""):
            ck = make_calibration_match_key(name, feature, unit, project_id=pid)
            if not ck or ck.startswith("vent_band|"):
                continue
            hit = by_cal.get(ck)
            if hit and hit.usable():
                return hit
    if nn in _NO_NAME_UNIT_FALLBACK:
        return None
    return by_name_unit.get((nn, un))


def _judge_cost(
    name: str,
    feature: str,
    unit: str,
    engine: PricingEngine,
    repo: KnowledgeRepository,
    ctx: PricingContext,
) -> CostParts:
    j = judge_line_components(name, feature, unit, engine, repo, ctx=ctx)
    return CostParts(
        material_main=float(j.material_main or 0),
        material_loss_rate=float(j.material_loss_rate or 0),
        labor=float(j.labor or 0),
        material_aux=float(j.material_aux or 0),
        machinery=float(j.machinery or 0),
        source=j.source or "judge",
    )


def _cost_parts_to_components(cost: CostParts) -> dict:
    return {
        "material_main": cost.material_main,
        "material_loss_rate": cost.material_loss_rate,
        "labor": cost.labor,
        "material_aux": cost.material_aux,
        "machinery": cost.machinery,
    }


def _fill_sheet_append(
    ws,
    layout,
    df,
    *,
    sn: str,
    by_key,
    by_name_unit,
    by_cal,
    engine,
    repo,
    ctx,
    wb_masters: Dict[str, MasterRowRef],
    quote_masters: Dict[str, MasterRowRef],
    rate_anchor_sheet: Optional[str] = None,
    sheet_bindings: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, int]:
    """任意甲方表：智能模板追加统一成本分析区。"""
    from src.export.cost_analysis_template import (
        bind_cost_analysis_template,
        bindings_summary,
        link_quote_unit_price,
        write_cost_headers,
        write_master_cost_row,
        write_rate_cell,
        write_slave_cost_row,
    )

    bindings = bind_cost_analysis_template(ws, layout)
    write_cost_headers(ws, layout, bindings)
    write_rate_cell(ws, bindings, rate_anchor_sheet=rate_anchor_sheet)
    if sheet_bindings is not None:
        sheet_bindings.append({"sheet": sn, **bindings_summary(bindings)})

    cm = layout.column_map
    name_i = cm.get("name")
    feat_i = cm.get("feature")
    unit_i = cm.get("unit")
    qty_i = cm.get("quantity")
    sh_ref = sh_judge = sh_master = sh_slave = sh_cross = 0

    for ri in range(layout.data_start_row, len(df)):
        row = tuple(df.iloc[ri])
        if not is_pricing_row(row, cm):
            continue
        excel_row = ri + 1
        name = _row_text(row, name_i)
        feature = _row_text(row, feat_i)
        unit = _row_text(row, unit_i)
        nn, un, sig = make_dedupe_key(name, feature, unit)
        key = f"{nn}|{sig}|{un}"

        if key not in wb_masters:
            wb_masters[key] = MasterRowRef(sn, excel_row)
            sh_master += 1
            ref = _lookup_reference(
                name, feature, unit, by_key, by_name_unit, by_cal
            )
            if ref and ref.usable():
                components = _cost_parts_to_components(ref)
                note = ref.source or ""
                sh_ref += 1
            else:
                cost = _judge_cost(name, feature, unit, engine, repo, ctx)
                components = _cost_parts_to_components(cost)
                note = cost.source or ""
                sh_judge += 1
            skip_cost = "合入他项" in note or not any(
                float(components.get(k) or 0)
                for k in (
                    "material_main",
                    "labor",
                    "material_aux",
                    "machinery",
                    "material_loss_rate",
                )
            )
            if skip_cost:
                components = {
                    "material_main": 0.0,
                    "material_loss_rate": 0.0,
                    "labor": 0.0,
                    "material_aux": 0.0,
                    "machinery": 0.0,
                }
            write_master_cost_row(
                ws,
                excel_row,
                bindings,
                components,
                blank_cost=skip_cost,
            )
            if not skip_cost and "防水" in name and "第二道" in name:
                quote_masters["js_waterproof"] = MasterRowRef(sn, excel_row)
            if skip_cost and "防水" in name and "js_waterproof" in quote_masters:
                link_quote_unit_price(
                    ws,
                    excel_row,
                    quote_masters["js_waterproof"],
                    bindings,
                )
        else:
            master = wb_masters[key]
            write_slave_cost_row(ws, excel_row, master, bindings)
            sh_slave += 1
            if master.sheet != sn:
                sh_cross += 1

    return {
        "masters": sh_master,
        "slaves": sh_slave,
        "cross_sheet_slaves": sh_cross,
        "reference": sh_ref,
        "judged": sh_judge,
        "template": bindings_summary(bindings),
    }


def _excel_copy_preserve(src: Path, dst: Path) -> None:
    """用 Excel 转换/复制，保留合并单元格、列宽等原表格式。"""
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb = None
    try:
        wb = excel.Workbooks.Open(str(src.resolve()))
        fmt = 51 if dst.suffix.lower() == ".xlsx" else 56
        dst.parent.mkdir(parents=True, exist_ok=True)
        if dst.exists():
            dst.unlink()
        wb.SaveAs(str(dst.resolve()), FileFormat=fmt)
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        excel.Quit()


def _materialize_workbook_copy(src: Path, out: Path) -> None:
    """复制招标工作簿；优先 Excel COM 保留格式，失败再 pandas 兜底。"""
    out.parent.mkdir(parents=True, exist_ok=True)
    src = src.resolve()
    out = out.resolve()
    if src.suffix.lower() == ".xls" or (
        out.suffix.lower() == ".xlsx" and src.suffix.lower() != ".xlsx"
    ):
        tmp = out.parent / f"{out.stem}_tmp{out.suffix}"
        if tmp.exists():
            tmp.unlink()
        try:
            _excel_copy_preserve(src, tmp)
            if out.exists():
                out.unlink()
            shutil.move(str(tmp), str(out))
            return
        except Exception:
            if tmp.exists():
                tmp.unlink()
        with pd.ExcelFile(src) as xl:
            with pd.ExcelWriter(out, engine="openpyxl") as writer:
                for sheet in xl.sheet_names:
                    df = pd.read_excel(xl, sheet_name=sheet, header=None)
                    df.to_excel(
                        writer,
                        sheet_name=sheet[:31],
                        index=False,
                        header=False,
                    )
        return
    if src != out:
        shutil.copy2(src, out)


def fill_inline_cost_link(
    tender_path: str | Path,
    reference_paths: List[str | Path],
    *,
    sheet_names: Optional[List[str]] = None,
    output_path: Optional[str | Path] = None,
    city: str = "珠海",
    tier: str = "mid",
    project_materials_ref: Optional[str] = None,
    db_path: Optional[str] = None,
) -> InlineLinkReport:
    src = Path(tender_path)
    out = Path(output_path) if output_path else (
        resolve_db_path(db_path).parent / "exports" / f"{src.stem}_成本链接.xlsx"
    )
    _materialize_workbook_copy(src, out)

    by_key, by_name_unit, by_cal = build_reference_price_index(reference_paths)
    repo = KnowledgeRepository(db_path)
    engine = PricingEngine(db_path)
    ctx = PricingContext(city=city, price_tier=tier)
    if project_materials_ref:
        ctx.project_materials_ref = project_materials_ref

    wb = load_workbook(out)
    target_sheets = sheet_names or [
        s for s in wb.sheetnames if not any(k in s for k in SKIP_SHEET_KEYWORDS)
    ]

    report = InlineLinkReport(output_path=str(out.resolve()))
    ref_hits = judged = masters = slaves = cross_sheet_slaves = 0
    wb_masters: Dict[str, MasterRowRef] = {}
    quote_masters: Dict[str, MasterRowRef] = {}
    rate_anchor_sheet: Optional[str] = None

    for sn in target_sheets:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        df = pd.read_excel(out, sheet_name=sn, header=None)
        layout = detect_sheet_layout(df, sn)
        if not layout:
            continue
        # 三种连接画像统一走智能模板（不写死 F/G/P，不新增说明表）
        stats = _fill_sheet_append(
            ws,
            layout,
            df,
            sn=sn,
            by_key=by_key,
            by_name_unit=by_name_unit,
            by_cal=by_cal,
            engine=engine,
            repo=repo,
            ctx=ctx,
            wb_masters=wb_masters,
            quote_masters=quote_masters,
            rate_anchor_sheet=rate_anchor_sheet,
        )
        profile = (stats.get("template") or {}).get("连接画像") or "append"
        if rate_anchor_sheet is None and profile == "national_quote":
            rate_anchor_sheet = sn

        ref_hits += stats["reference"]
        judged += stats["judged"]
        masters += stats["masters"]
        slaves += stats["slaves"]
        cross_sheet_slaves += stats["cross_sheet_slaves"]
        report.sheets.append({"sheet": sn, "layout": profile, **stats})

    wb.save(out)
    report.reference_hits = ref_hits
    report.judged = judged
    report.masters = masters
    report.slaves = slaves
    report.cross_sheet_slaves = cross_sheet_slaves
    return report
