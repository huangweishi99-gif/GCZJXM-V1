# -*- coding: utf-8 -*-
"""临时：探索三份云湾文件结构（表头/列含义/控制价列），用后即删。"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import win32com.client
from openpyxl.utils import get_column_letter

XLS_A = r"清单数据资料\甲方招标清单\云湾花园项目-5#6#户内、公区，样板，架空层装修工程专业分包-投标清单-20260825.xls"
XLS_B = r"清单数据资料\甲方招标清单\云湾花园项目-会所售楼部装修工程专业分包-投标清单-20260825.xls"
XLSX = r"清单数据资料\甲方招标清单\会所、5#6#楼公区及户内、样板间、架空层安装工程投标清单-20260825.xlsx"

xl = win32com.client.DispatchEx("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    for tag, path, sheet_name in (("5#6#", XLS_A, "清单表"), ("会所xls", XLS_B, "清单表")):
        print("\n##########", tag)
        wb = xl.Workbooks.Open(os.path.abspath(path), UpdateLinks=0, ReadOnly=True)
        ws = wb.Worksheets(sheet_name)
        ur = ws.UsedRange
        print("UsedRange:", ur.Row, ur.Column, ur.Rows.Count, ur.Columns.Count)
        # 打印前 8 行所有列
        for r in range(1, 9):
            vals = []
            for c in range(1, ur.Columns.Count + 1):
                v = ws.Cells(r, c).Value
                if v is not None:
                    vals.append(f"{get_column_letter(c)}{r}={v!r}")
            print(" | ".join(vals))
        wb.Close(False)
finally:
    xl.Quit()

print("\n########## 会所安装xlsx")
import openpyxl
wb = openpyxl.load_workbook(XLSX, data_only=False)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"-- sheet[{sn}] dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
wb.close()
# 消防电/消防水 详细表头
wb = openpyxl.load_workbook(XLSX, data_only=True)
for sn in ("会所售楼处消防电工程", "会所售楼处消防水工程"):
    ws = wb[sn]
    print(f"\n## {sn} 前10行")
    for r in range(1, 11):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                vals.append(f"{get_column_letter(c)}{r}={v!r}")
        if vals:
            print(" | ".join(vals))
wb.close()
