# -*- coding: utf-8 -*-
"""解析 Excel 单元格公式（工程量列常用，含跨表引用）。"""
from __future__ import annotations

import re
from typing import Any, Dict, Optional, Tuple

from openpyxl.utils import column_index_from_string

_CELL_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_CROSS_SHEET_REF = re.compile(
    r"(?:'([^']+)'|([^!'\"&]+))!(\$?)([A-Z]{1,3})(\$?)(\d+)",
    re.I,
)


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None


def _resolve_sheet_name(wb, raw: str) -> Optional[str]:
    name = (raw or "").strip().strip("'\"")
    if not name or wb is None:
        return None
    if name in wb.sheetnames:
        return name
    for sn in wb.sheetnames:
        if name in sn or sn in name:
            return sn
    return None


def _eval_cross_sheet(
    wb,
    sheet_raw: str,
    col_letters: str,
    row_num: int,
    cache: Optional[Dict[Tuple[str, int, int], Optional[float]]],
    depth: int,
) -> Optional[float]:
    sn = _resolve_sheet_name(wb, sheet_raw)
    if not sn:
        return None
    ws = wb[sn]
    col = column_index_from_string(col_letters.upper())
    return eval_cell(ws, int(row_num), col, cache, depth, wb=wb)


def eval_cell(
    ws,
    row: int,
    col: int,
    cache: Optional[Dict[Tuple[Any, ...], Optional[float]]] = None,
    depth: int = 0,
    *,
    wb=None,
) -> Optional[float]:
    """求值 1-based 单元格；支持 =A1+B1、=汇总表!C16 等。"""
    if depth > 16:
        return None
    book = wb or getattr(ws, "parent", None)
    key = (id(ws), row, col)
    if cache is not None:
        if key in cache:
            return cache[key]
    val = ws.cell(row, col).value
    if val is None:
        if cache is not None:
            cache[key] = None
        return None
    if isinstance(val, (int, float)):
        out = float(val)
        if cache is not None:
            cache[key] = out
        return out
    s = str(val).strip()
    if not s.startswith("="):
        out = _to_float(s)
        if cache is not None:
            cache[key] = out
        return out
    out = _eval_expr(ws, s[1:].replace(" ", ""), row, col, cache, depth + 1, book)
    if cache is not None:
        cache[key] = out
    return out


def _eval_expr(
    ws,
    expr: str,
    row: int,
    col: int,
    cache: Optional[Dict[Tuple[Any, ...], Optional[float]]],
    depth: int,
    wb,
) -> Optional[float]:
    if not expr:
        return None

    solo = _CROSS_SHEET_REF.fullmatch(expr)
    if solo:
        sheet = solo.group(1) or solo.group(2) or ""
        return _eval_cross_sheet(wb, sheet, solo.group(4), int(solo.group(6)), cache, depth)

    def repl_cross(m: re.Match) -> str:
        sheet = m.group(1) or m.group(2) or ""
        v = _eval_cross_sheet(wb, sheet, m.group(4), int(m.group(6)), cache, depth)
        return str(v if v is not None else 0)

    expr2 = _CROSS_SHEET_REF.sub(repl_cross, expr)
    if "nan" in expr2:
        return None

    def repl(m: re.Match) -> str:
        c = column_index_from_string(m.group(1))
        r = int(m.group(2))
        v = eval_cell(ws, r, c, cache, depth, wb=wb)
        return str(v if v is not None else 0)

    replaced = _CELL_REF.sub(repl, expr2)
    if "nan" in replaced:
        return None
    if not re.match(r"^[\d.+\-*/()]+$", replaced):
        return None
    try:
        return float(eval(replaced))  # noqa: S307 — trusted local xlsx
    except Exception:
        return None


def row_quantity_from_ws(
    ws,
    row: int,
    column_map: dict,
    floor_cols=None,
) -> Optional[float]:
    """结合公式求值读取合计/分层工程量。"""
    cache: Dict[Tuple[Any, ...], Optional[float]] = {}
    wb = getattr(ws, "parent", None)
    qty_i = column_map.get("quantity")
    if qty_i is not None:
        v = eval_cell(ws, row, qty_i + 1, cache, wb=wb)
        if v is not None:
            return v
    if floor_cols:
        total = 0.0
        found = False
        for idx, _ in floor_cols:
            v = eval_cell(ws, row, idx + 1, cache, wb=wb)
            if v is not None:
                total += v
                found = True
        return total if found else None
    return None
