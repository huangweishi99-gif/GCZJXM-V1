# -*- coding: utf-8 -*-
"""
成本分析智能模板 —— 统一成本区 + 按甲方表自动连接。

闭环金标准归纳出三种「连接画像」（成本区结构永远相同）：

1. haide_inquiry（清单询，如售楼处/地库/会所合并）
   - 甲方左区已有：人工 / 主材 / 辅材及机械费 / 管理费 / 综合单价 / 合价
   - 成本区接在表尾后（通常 P 起）：成本单价、主材、损耗率、人工、辅材、机械、合价
   - 左区用公式桥接（不填死数）：
     人工 = 成本人工 × $人工$1
     主材 = 成本主材×(1+损耗) × $主材$1
     辅材及机械 = (辅材+机械) × $辅材$1
   - 费率默认写在成本主材 / 人工 / 辅材列第 1 行

2. tender_breakdown（国标「其中」分解，如开市客）
   - 甲方左区：综合单价 + 其中(人工费/材料费/其他费) + 合价
   - 成本区仍是同一套 7 列，接在甲方末列后空 1 列
   - 桥接：
     人工费 = 成本人工 × $人工$1
     材料费 = 主材(1+损耗)×$主材$1 + 辅材×$辅材$1 + 机械
     其他费 = 成本单价 × $成本单价$1（常为 0）
     综合单价 = 人工费+材料费+其他费
     合价 = 工程量 × 综合单价

3. national_quote（国标简表，如 B2 户型）
   - 甲方左区：综合单价 + 合价（无「其中」分解）
   - 成本区同一套 7 列追加
   - 桥接：综合单价 = 成本单价 × $成本单价$1；合价 = 工程量 × 综合单价

铁律：
- 甲方原有表头 / 样式 / 数据区结构一律不许改写为死数；
- 只在甲方末列之后写成本分析，并按画像用公式桥接甲方报价列；
- 成本合价 = 成本单价 × 工程量。
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MASTER_FILL = PatternFill("solid", fgColor="C6EFCE")
SLAVE_FILL = PatternFill("solid", fgColor="FFF2CC")

from src.ingest.detector import MAX_REASONABLE_COL, SheetLayout
from src.link.dedupe import MasterRowRef, master_cell_formula

# 三种连接画像（稳定标识，勿轻易改名）
PROFILE_HAIDE = "haide_inquiry"
PROFILE_BREAKDOWN = "tender_breakdown"
PROFILE_NATIONAL = "national_quote"

def _existing_cost_positions(m: dict):
    """表头已含成本区（detector 标记 _cost_side / *_cost）时，复用其 1-based 列号。

    避免在已带成本区的文件（如已填价/金标准）上重复追加成本列，导致错位覆盖。
    优先 *_cost 字段；缺失则按 cost_unit_price 锚定顺序补齐。
    """
    if not m.get("_cost_side"):
        return None
    anchor = m.get("cost_unit_price")
    if not isinstance(anchor, int):
        return None
    start = anchor + 1
    positions: dict = {}
    for i, f in enumerate(COST_FIELDS):
        key = f
        if f in (
            "material_main",
            "material_loss_rate",
            "labor",
            "material_aux",
            "machinery",
        ):
            ck = f + "_cost"
            if isinstance(m.get(ck), int):
                key = ck
        v = m.get(key)
        positions[f] = v + 1 if isinstance(v, int) else start + i
    return positions


# 固定成本分析列（顺序不可变）
COST_FIELDS: Tuple[str, ...] = (
    "cost_unit_price",
    "material_main",
    "material_loss_rate",
    "labor",
    "material_aux",
    "machinery",
    "cost_amount",
)

COST_HEADERS: Tuple[Tuple[str, str], ...] = (
    ("cost_unit_price", "成本单价"),
    ("material_main", "主材"),
    ("material_loss_rate", "损耗率"),
    ("labor", "人工"),
    ("material_aux", "辅材"),
    ("machinery", "机械"),
    ("cost_amount", "合价"),
)

# 甲方末列后空几列再写成本区：1 →「第 2 列起」
DEFAULT_GAP_COLS = 1


@dataclass
class CostAnalysisBindings:
    """一张 sheet 上甲方列 + 成本分析列的绑定结果（均为 1-based 列号）。"""

    profile: str
    client_last_col: int
    cost_start_col: int
    # 成本区
    cost_unit_price: int
    material_main: int
    material_loss_rate: int
    labor: int
    material_aux: int
    machinery: int
    cost_amount: int
    # 甲方识别列
    quantity: Optional[int] = None
    unit_price: Optional[int] = None
    amount: Optional[int] = None
    name: Optional[int] = None
    feature: Optional[int] = None
    unit: Optional[int] = None
    remark: Optional[int] = None
    # 甲方「其中」分解列（tender_breakdown / haide）
    client_labor: Optional[int] = None
    client_main: Optional[int] = None
    client_aux_mach: Optional[int] = None
    client_other: Optional[int] = None
    # 行为开关
    quote_from_cost: bool = True
    gap_cols: int = DEFAULT_GAP_COLS
    # 表头已含成本区（已填价/金标准/甲方自带成本列）：复用列且不重写表头
    has_existing_cost: bool = False

    def col(self, field: str) -> Optional[int]:
        return getattr(self, field, None)

    def letter(self, field: str) -> Optional[str]:
        c = self.col(field)
        return get_column_letter(c) if c else None

    def cost_field_cols(self) -> List[int]:
        return [self.col(f) for f in COST_FIELDS if self.col(f)]


def _header_blob(ws, layout: SheetLayout, max_col: int = 16) -> str:
    rows = [layout.header_row + 1]
    if layout.subheader_row is not None:
        rows.append(layout.subheader_row + 1)
    else:
        rows.append(layout.header_row + 2)
    parts: List[str] = []
    for r in rows:
        if r > ws.max_row:
            continue
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                parts.append(str(v).replace(" ", "").replace("\n", ""))
    return "".join(parts)


def _find_label_col(ws, layout: SheetLayout, labels: Sequence[str], max_col: int = 16) -> Optional[int]:
    rows = [layout.header_row + 1]
    if layout.subheader_row is not None:
        rows.append(layout.subheader_row + 1)
    else:
        rows.append(layout.header_row + 2)
    for r in rows:
        if r > ws.max_row:
            continue
        for c in range(1, max_col + 1):
            v = str(ws.cell(r, c).value or "").replace(" ", "").replace("\n", "")
            if not v:
                continue
            for lab in labels:
                if lab in v:
                    return c
    return None


def detect_connection_profile(ws, layout: SheetLayout) -> str:
    """根据甲方原表头识别连接画像（不依赖固定列号）。"""
    blob = _header_blob(ws, layout)
    m = layout.column_map

    # 开市客等：综合单价 + 其中(人工费/材料费/其他费)
    if (
        "人工费" in blob
        and "材料费" in blob
        and ("其他费" in blob or "其它费" in blob)
    ):
        return PROFILE_BREAKDOWN

    # 清单询：左区已有人工/主材/辅材及机械费，在综合单价左侧
    labor_i = m.get("labor")
    main_i = m.get("material_main")
    up_i = m.get("unit_price")
    if (
        isinstance(labor_i, int)
        and isinstance(main_i, int)
        and isinstance(up_i, int)
        and labor_i < up_i
        and main_i < up_i
        and "主材" in blob
        and "人工费" not in blob
    ):
        return PROFILE_HAIDE

    # 兼容：表头写了「人工」「主材」「辅材及机械」但 detector 列名略偏
    if (
        "人工" in blob
        and "主材" in blob
        and ("辅材及机械" in blob or "辅材机械" in blob)
        and "综合单价" in blob
        and "人工费" not in blob
    ):
        return PROFILE_HAIDE

    return PROFILE_NATIONAL


def is_haide_inquiry_sheet(ws, layout: SheetLayout) -> bool:
    return detect_connection_profile(ws, layout) == PROFILE_HAIDE


def detect_delivery_mode(ws, layout: SheetLayout) -> str:
    """统一走智能模板；返回连接画像标识供报告展示。"""
    return detect_connection_profile(ws, layout)


def effective_client_last_column(ws, layout: SheetLayout) -> int:
    """甲方区域末列：以备注/合价等业务列为主，最多再右扩 2 列（如「检查」）。"""
    m = layout.column_map
    anchors: List[int] = []
    for key in (
        "remark",
        "amount",
        "unit_price",
        "management",
        "profit",
        "machinery",
        "labor",
        "material_main",
        "quantity",
        "unit",
        "feature",
        "name",
    ):
        idx = m.get(key)
        # 金标准已带成本侧时，人工/主材可能是成本列，只用报价语义列锚定
        if isinstance(idx, int):
            if m.get("_cost_side") and key in (
                "labor",
                "material_main",
                "machinery",
                "material_aux",
                "material_loss_rate",
            ):
                continue
            anchors.append(idx + 1)
    base = max(anchors) if anchors else 1

    last = base
    rows = [layout.header_row + 1]
    if layout.subheader_row is not None:
        rows.append(layout.subheader_row + 1)
    elif layout.header_row + 2 <= ws.max_row:
        rows.append(layout.header_row + 2)
    # 只在 base 右侧最多探 2 列，避免把零散检查列/图档列当成甲方区
    for r in rows:
        for c in range(base, min(base + 2, MAX_REASONABLE_COL) + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                last = max(last, c)
    return max(last, base, 1)


def bind_cost_analysis_template(
    ws,
    layout: SheetLayout,
    *,
    gap_cols: int = DEFAULT_GAP_COLS,
    profile: Optional[str] = None,
) -> CostAnalysisBindings:
    """根据表头智能识别甲方列 + 连接画像，在末列后绑定固定成本分析区。"""
    profile = profile or detect_connection_profile(ws, layout)
    m = layout.column_map
    client_last = effective_client_last_column(ws, layout)
    # 表头已含成本区（已填价/金标准/甲方自带成本列）：复用已有列，避免重复追加破坏
    positions = _existing_cost_positions(m)
    if positions is None:
        # 清单询历史口径从 P 列起；若甲方已伸到 N 之后则跟随末列 + gap
        if profile == PROFILE_HAIDE:
            start = max(16, client_last + max(gap_cols, 0) + 1)
        else:
            start = client_last + max(gap_cols, 0) + 1
        positions = {f: start + i for i, f in enumerate(COST_FIELDS)}

    def _cli(field: str) -> Optional[int]:
        idx = m.get(field)
        return idx + 1 if isinstance(idx, int) else None

    # 甲方「其中」列（画像相关）
    client_labor = client_main = client_aux_mach = client_other = None
    if profile == PROFILE_HAIDE:
        client_labor = _cli("labor") or _find_label_col(ws, layout, ("人工",))
        client_main = _cli("material_main") or _find_label_col(ws, layout, ("主材",))
        client_aux_mach = _cli("machinery") or _find_label_col(
            ws, layout, ("辅材及机械费", "辅材及机械", "辅材机械")
        )
    elif profile == PROFILE_BREAKDOWN:
        client_labor = _find_label_col(ws, layout, ("人工费",))
        client_main = _find_label_col(ws, layout, ("材料费",))
        client_other = _find_label_col(ws, layout, ("其他费", "其它费"))

    return CostAnalysisBindings(
        profile=profile,
        client_last_col=client_last,
        cost_start_col=positions["cost_unit_price"],
        cost_unit_price=positions["cost_unit_price"],
        material_main=positions["material_main"],
        material_loss_rate=positions["material_loss_rate"],
        labor=positions["labor"],
        material_aux=positions["material_aux"],
        machinery=positions["machinery"],
        cost_amount=positions["cost_amount"],
        quantity=_cli("quantity"),
        unit_price=_cli("unit_price"),
        amount=_cli("amount"),
        name=_cli("name"),
        feature=_cli("feature"),
        unit=_cli("unit"),
        remark=_cli("remark"),
        client_labor=client_labor,
        client_main=client_main,
        client_aux_mach=client_aux_mach,
        client_other=client_other,
        quote_from_cost=True,
        gap_cols=gap_cols,
        has_existing_cost=bool(m.get("_cost_side")),
    )


def _safe_set(ws, row: int, col: Optional[int], value) -> None:
    if not col:
        return
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def _safe_ref(ws, row: int, col: Optional[int], formula: str) -> None:
    _safe_set(ws, row, col, formula)


def write_cost_headers(ws, layout: SheetLayout, b: CostAnalysisBindings) -> None:
    """仅在成本追加列写统一表头。

    - 有子表头（国标简表等）：主表头写「成本单价分析」合并，子表头写 7 列名
    - 无子表头（清单询 / 开市客）：7 列名直接写在甲方主表头行，避免落到数据行
    - 表头已含成本区（has_existing_cost）：原表头完整，直接保留
    """
    if b.has_existing_cost:
        return
    hr = layout.header_row + 1
    has_sub = layout.subheader_row is not None
    title_row = hr if has_sub else None
    field_row = (layout.subheader_row + 1) if has_sub else hr

    if title_row:
        c0, c1 = b.cost_unit_price, b.cost_amount
        cell = ws.cell(title_row, c0)
        if not isinstance(cell, MergedCell):
            cell.value = "成本单价分析"
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        if c1 > c0:
            try:
                ws.merge_cells(
                    start_row=title_row,
                    start_column=c0,
                    end_row=title_row,
                    end_column=c1,
                )
            except ValueError:
                pass

    for field, title in COST_HEADERS:
        c = b.col(field)
        if not c:
            continue
        sub = ws.cell(field_row, c)
        if isinstance(sub, MergedCell):
            continue
        sub.value = title
        sub.font = Font(bold=True)
        sub.alignment = Alignment(horizontal="center", wrap_text=True)


def _rate_formula(ws, col: int, value, rate_anchor_sheet: Optional[str]) -> None:
    letter = get_column_letter(col)
    if not rate_anchor_sheet or rate_anchor_sheet == ws.title:
        _safe_set(ws, 1, col, value)
    else:
        _safe_set(ws, 1, col, f"='{rate_anchor_sheet}'!${letter}$1")


def write_rate_cell(
    ws,
    b: CostAnalysisBindings,
    *,
    rate_anchor_sheet: Optional[str] = None,
) -> None:
    """按画像写第 1 行费率系数。

    - 清单询 / 其中分解：每张表本地写系数（金标准）
    - 国标简表：仅成本单价列；非首表引用首表费用率
    """
    if b.profile == PROFILE_HAIDE:
        _rate_formula(ws, b.material_main, 1, None)
        _rate_formula(ws, b.labor, 1, None)
        _rate_formula(ws, b.material_aux, 1, None)
        return
    if b.profile == PROFILE_BREAKDOWN:
        _rate_formula(ws, b.cost_unit_price, 0, None)
        _rate_formula(ws, b.material_main, 1, None)
        _rate_formula(ws, b.labor, 1, None)
        _rate_formula(ws, b.material_aux, 1, None)
        return
    # national_quote：仅成本单价列费率 1，跨表引用首表
    _rate_formula(ws, b.cost_unit_price, 1, rate_anchor_sheet)


def write_cost_formulas(ws, row: int, b: CostAnalysisBindings) -> None:
    """成本单价 = 主材+主材×损耗+人工+辅材+机械；成本合价 = 成本单价×工程量。"""
    main_l = b.letter("material_main")
    loss_l = b.letter("material_loss_rate")
    labor_l = b.letter("labor")
    aux_l = b.letter("material_aux")
    mach_l = b.letter("machinery")
    cost_l = b.letter("cost_unit_price")
    qty_l = b.letter("quantity")
    amt_l = b.letter("cost_amount")
    if main_l and loss_l and labor_l and aux_l and mach_l and cost_l:
        _safe_ref(
            ws,
            row,
            b.cost_unit_price,
            (
                f"={main_l}{row}+{main_l}{row}*{loss_l}{row}"
                f"+{labor_l}{row}+{aux_l}{row}+{mach_l}{row}"
            ),
        )
    if cost_l and qty_l and amt_l:
        _safe_ref(
            ws,
            row,
            b.cost_amount,
            f"={cost_l}{row}*{qty_l}{row}",
        )


def write_quote_bridge(ws, row: int, b: CostAnalysisBindings) -> None:
    """按画像把成本区公式桥接到甲方报价列（只写公式，不填死数）。"""
    if not b.quote_from_cost:
        return

    main_l = b.letter("material_main")
    loss_l = b.letter("material_loss_rate")
    labor_l = b.letter("labor")
    aux_l = b.letter("material_aux")
    mach_l = b.letter("machinery")
    cost_l = b.letter("cost_unit_price")
    qty_l = b.letter("quantity")
    up_l = b.letter("unit_price")
    amt_l = b.letter("amount")

    if b.profile == PROFILE_HAIDE:
        # F/G/H ← S/Q/T+U
        if b.client_labor and labor_l:
            _safe_ref(ws, row, b.client_labor, f"={labor_l}{row}*${labor_l}$1")
        if b.client_main and main_l and loss_l:
            _safe_ref(
                ws,
                row,
                b.client_main,
                f"={main_l}{row}*(1+{loss_l}{row})*${main_l}$1",
            )
        if b.client_aux_mach and aux_l and mach_l:
            _safe_ref(
                ws,
                row,
                b.client_aux_mach,
                f"=({aux_l}{row}+{mach_l}{row})*${aux_l}$1",
            )
        return

    if b.profile == PROFILE_BREAKDOWN:
        # 人工费 / 材料费 / 其他费 / 综合单价 / 合价
        if b.client_labor and labor_l:
            _safe_ref(ws, row, b.client_labor, f"={labor_l}{row}*${labor_l}$1")
        if b.client_main and main_l and loss_l and aux_l and mach_l:
            _safe_ref(
                ws,
                row,
                b.client_main,
                (
                    f"=({main_l}{row}+{main_l}{row}*{loss_l}{row})*${main_l}$1"
                    f"+{aux_l}{row}*${aux_l}$1+{mach_l}{row}"
                ),
            )
        if b.client_other and cost_l:
            _safe_ref(ws, row, b.client_other, f"={cost_l}{row}*${cost_l}$1")
        # 综合单价 = 三者之和（按列存在情况拼接）
        parts = []
        for c in (b.client_labor, b.client_main, b.client_other):
            if c:
                parts.append(f"{get_column_letter(c)}{row}")
        if parts and b.unit_price:
            _safe_ref(ws, row, b.unit_price, "=" + "+".join(parts))
        if qty_l and up_l and b.amount:
            _safe_ref(ws, row, b.amount, f"={up_l}{row}*{qty_l}{row}")
        return

    # national_quote
    if cost_l and b.unit_price:
        _safe_ref(ws, row, b.unit_price, f"={cost_l}{row}*${cost_l}$1")
    if qty_l and up_l and b.amount:
        _safe_ref(ws, row, b.amount, f"={qty_l}{row}*{up_l}{row}")


def _apply_cost_fill(ws, row: int, b: CostAnalysisBindings, fill: PatternFill) -> None:
    """清单询金标准：母行浅绿、链接行浅黄（仅成本人材机列）。"""
    if b.profile != PROFILE_HAIDE:
        return
    for field in (
        "material_main",
        "material_loss_rate",
        "labor",
        "material_aux",
        "machinery",
    ):
        c = b.col(field)
        if not c:
            continue
        cell = ws.cell(row, c)
        if not isinstance(cell, MergedCell):
            cell.fill = fill


def write_master_cost_row(
    ws,
    row: int,
    b: CostAnalysisBindings,
    components: Dict[str, float],
    *,
    blank_cost: bool = False,
) -> None:
    """母行：人材机数值 + 成本公式 + 报价桥接。"""
    if not blank_cost:
        main = float(components.get("material_main") or 0)
        loss = float(components.get("material_loss_rate") or 0)
        labor = float(components.get("labor") or 0)
        aux = float(components.get("material_aux") or 0)
        mach = float(components.get("machinery") or 0)
        _safe_set(ws, row, b.material_main, round(main, 4) if main else None)
        _safe_set(ws, row, b.material_loss_rate, loss)
        _safe_set(ws, row, b.labor, round(labor, 4) if labor else None)
        _safe_set(ws, row, b.material_aux, round(aux, 4) if aux else None)
        _safe_set(ws, row, b.machinery, round(mach, 4) if mach else None)
    write_cost_formulas(ws, row, b)
    write_quote_bridge(ws, row, b)
    _apply_cost_fill(ws, row, b, MASTER_FILL)


def write_slave_cost_row(
    ws,
    row: int,
    master: MasterRowRef,
    b: CostAnalysisBindings,
) -> None:
    """重复项：人材机链母行；清单询左区 F/G/H 直接链母行，其余画像用本行桥接公式。"""
    local = ws.title
    for field in (
        "material_main",
        "material_loss_rate",
        "labor",
        "material_aux",
        "machinery",
    ):
        c = b.col(field)
        letter = b.letter(field)
        if c and letter:
            _safe_set(ws, row, c, master_cell_formula(master, letter, local))
    write_cost_formulas(ws, row, b)
    if b.profile == PROFILE_HAIDE:
        # 金标准：子行左区直接链母行 F/G/H（跨 sheet 写全表名）
        for client_col in (b.client_labor, b.client_main, b.client_aux_mach):
            if not client_col:
                continue
            letter = get_column_letter(client_col)
            _safe_set(ws, row, client_col, master_cell_formula(master, letter, local))
    else:
        write_quote_bridge(ws, row, b)
    _apply_cost_fill(ws, row, b, SLAVE_FILL)


def link_quote_unit_price(
    ws,
    row: int,
    master: MasterRowRef,
    b: CostAnalysisBindings,
) -> None:
    """合入他项：综合单价链接母行，合价=工程量×综合单价。"""
    up_c = b.unit_price
    up_l = b.letter("unit_price")
    qty_l = b.letter("quantity")
    amt_c = b.amount
    amt_l = b.letter("amount")
    if up_c and up_l:
        _safe_set(ws, row, up_c, master_cell_formula(master, up_l, ws.title))
    if qty_l and up_l and amt_c and amt_l:
        _safe_ref(ws, row, amt_c, f"={qty_l}{row}*{up_l}{row}")


def bindings_summary(b: CostAnalysisBindings) -> Dict[str, object]:
    return {
        "连接画像": b.profile,
        "甲方末列": b.client_last_col,
        "成本起始列": b.cost_start_col,
        "空列间隔": b.gap_cols,
        "工程量列": b.quantity,
        "综合单价列": b.unit_price,
        "投标合价列": b.amount,
        "甲方人工列": b.client_labor,
        "甲方主材/材料费列": b.client_main,
        "甲方其他费列": b.client_other,
        "成本单价列": b.cost_unit_price,
        "成本合价列": b.cost_amount,
    }
