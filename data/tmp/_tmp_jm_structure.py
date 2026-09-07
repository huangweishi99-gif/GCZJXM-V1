# -*- coding: utf-8 -*-
"""读取江门两个 Excel 的 sheet 与列结构"""
import io, sys
from pathlib import Path
sys.path.insert(0, r'd:\projects\cost-estimation')

import openpyxl

files = [
    r'd:\projects\cost-estimation\清单数据资料\AI学习清单\江门新会骏凯豪庭五期三标项目1栋05户型精装修工程-投标报价-20260818.xlsx',
    r'd:\projects\cost-estimation\data\exports\江门新会骏凯豪庭五期三标项目1栋05户型精装修工程-投标报价_成本填写.xlsx',
]

out = io.open(r'd:\projects\cost-estimation\data\tmp\_tmp_jm_structure_out.txt', 'w', encoding='utf-8')
for f in files:
    out.write(f'===== FILE: {Path(f).name} =====\n')
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    for ws in wb.worksheets:
        out.write(f'  --- sheet: {ws.title}  max_row={ws.max_row} max_col={ws.max_column}\n')
        # 打印前 8 行每行前 15 列
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, max_col=min(ws.max_column, 15), values_only=True)):
            vals = ['' if v is None else str(v)[:22] for v in row]
            out.write(f'    r{i+1}: {vals}\n')
        if ws.max_row > 8:
            out.write(f'    ... 共 {ws.max_row} 行\n')
        out.write('\n')
    wb.close()
out.close()
print('done')
