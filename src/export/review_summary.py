# -*- coding: utf-8 -*-
"""交付复核摘要：另存工作簿，不改甲方原表。"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


_HEADERS = (
    "工作表", "序号", "项目名称", "单位", "工程量",
    "状态", "来源", "影响合价", "说明", "建议定额",
)


def write_review_workbook(
    path: str | Path,
    rows: List[Dict[str, Any]],
    boq_report=None,
) -> Path:
    """未组价、体检拦截、定额建议和警告集中到一页，按影响合价降序。"""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "造价员复核"

    fatal = getattr(boq_report, "fatal_count", 0) or 0
    warning = getattr(boq_report, "warning_count", 0) or 0
    ws["A1"] = f"体检：致命 {fatal}，警告 {warning}。致命项未自动填价；定额命中只作建议。"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:J1")

    for col, title in enumerate(_HEADERS, 1):
        cell = ws.cell(2, col, title)
        cell.font = Font(bold=True)

    ordered = sorted(
        rows,
        key=lambda r: (
            0 if r.get("status") == "不自动填" else 1,
            -abs(float(r.get("amount_hint") or 0)),
        ),
    )
    for i, row in enumerate(ordered, 3):
        values = (
            row.get("sheet") or "",
            row.get("seq") or "",
            row.get("name") or "",
            row.get("unit") or "",
            row.get("quantity"),
            row.get("status") or "",
            row.get("source") or "",
            row.get("amount_hint") or 0,
            row.get("note") or "",
            row.get("quota_hint") or "",
        )
        for col, value in enumerate(values, 1):
            ws.cell(i, col, value)

    widths = (16, 10, 36, 8, 12, 14, 22, 14, 60, 48)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A2:J{max(2, len(ordered) + 2)}"
    ws.freeze_panes = "A3"
    wb.save(out)
    return out
